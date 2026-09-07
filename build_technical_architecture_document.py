from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT='Documentacion_Tecnica_y_Arquitectura.docx'
BLUE='2E74B5'; DARK='1F4D78'; INK='203040'; PALE='E8EEF5'; GRAY='F2F4F7'

def set_cell(cell, fill=None):
    tcPr=cell._tc.get_or_add_tcPr()
    if fill:
        shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),fill); tcPr.append(shd)
    mar=OxmlElement('w:tcMar')
    for side,val in [('top','80'),('start','120'),('bottom','80'),('end','120')]:
        e=OxmlElement('w:'+side); e.set(qn('w:w'),val); e.set(qn('w:type'),'dxa'); mar.append(e)
    tcPr.append(mar); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER

def geometry(t,widths):
    t.autofit=False; p=t._tbl.tblPr
    w=p.first_child_found_in('w:tblW'); w.set(qn('w:w'),str(sum(widths))); w.set(qn('w:type'),'dxa')
    ind=OxmlElement('w:tblInd'); ind.set(qn('w:w'),'120'); ind.set(qn('w:type'),'dxa'); p.append(ind)
    for col,n in zip(t._tbl.tblGrid.gridCol_lst,widths): col.set(qn('w:w'),str(n))
    for row in t.rows:
        for cell,n in zip(row.cells,widths):
            cell.width=Inches(n/1440); cell._tc.tcPr.tcW.set(qn('w:w'),str(n)); cell._tc.tcPr.tcW.set(qn('w:type'),'dxa'); set_cell(cell)

def table(doc, heads, rows, widths):
    t=doc.add_table(rows=1,cols=len(heads)); t.style='Table Grid'; t.alignment=WD_TABLE_ALIGNMENT.LEFT
    for c,txt in zip(t.rows[0].cells,heads):
        set_cell(c,PALE); p=c.paragraphs[0]; p.paragraph_format.space_after=Pt(0); r=p.add_run(txt); r.bold=True; r.font.size=Pt(9); r.font.color.rgb=RGBColor.from_string(DARK)
    for row in rows:
        cells=t.add_row().cells
        for c,txt in zip(cells,row):
            p=c.paragraphs[0]; p.paragraph_format.space_after=Pt(0); r=p.add_run(txt); r.font.size=Pt(9)
    geometry(t,widths); doc.add_paragraph().paragraph_format.space_after=Pt(1)

doc=Document(); sec=doc.sections[0]
sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Inches(1); sec.header_distance=sec.footer_distance=Inches(.492)
styles=doc.styles
styles['Normal'].font.name='Calibri'; styles['Normal'].font.size=Pt(11); styles['Normal'].font.color.rgb=RGBColor.from_string(INK); styles['Normal'].paragraph_format.space_after=Pt(6); styles['Normal'].paragraph_format.line_spacing=1.10
for n,size,color,before,after in [(1,16,BLUE,16,8),(2,13,BLUE,12,6),(3,12,DARK,8,4)]:
    st=styles[f'Heading {n}']; st.font.name='Calibri'; st.font.size=Pt(size); st.font.bold=True; st.font.color.rgb=RGBColor.from_string(color); st.paragraph_format.space_before=Pt(before); st.paragraph_format.space_after=Pt(after); st.paragraph_format.keep_with_next=True
head=sec.header.paragraphs[0]; head.add_run('PLATAFORMA DE REEMBOLSOS  |  ARQUITECTURA TÉCNICA').font.size=Pt(8); head.runs[0].font.color.rgb=RGBColor.from_string('6B7280')
foot=sec.footer.paragraphs[0]; foot.alignment=WD_ALIGN_PARAGRAPH.RIGHT; foot.add_run('Referencia técnica  |  3 de septiembre de 2026').font.size=Pt(8)

p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(42); p.paragraph_format.space_after=Pt(8); r=p.add_run('PLATAFORMA DE REEMBOLSOS'); r.bold=True; r.font.size=Pt(11); r.font.color.rgb=RGBColor.from_string(BLUE)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(10); r=p.add_run('Documentación técnica y arquitectura'); r.bold=True; r.font.size=Pt(24); r.font.color.rgb=RGBColor.from_string(DARK)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(22); r=p.add_run('Estado actual basado en el código fuente del repositorio'); r.italic=True; r.font.size=Pt(12); r.font.color.rgb=RGBColor.from_string('5B6573')
table(doc,['Control','Valor'], [('Versión','1.0'),('Fecha','3 de septiembre de 2026'),('Audiencia','Desarrollo, soporte, administración técnica y auditoría'),('Alcance','Aplicación Laravel web, persistencia, archivos, correo, seguridad y operación')],[2450,6910])

doc.add_heading('1. Resumen ejecutivo',1)
doc.add_paragraph('La plataforma es una aplicación web Laravel para gestionar reembolsos, comprobaciones, viáticos, fondos fijos y eventos. Su núcleo combina captura y validación de CFDI, flujos dinámicos de aprobación, colas de Cuentas por Pagar, evidencias almacenadas, notificaciones y auditoría.')
doc.add_paragraph('La arquitectura sigue el patrón MVC de Laravel: Blade/Tailwind/Alpine en presentación; rutas y controladores para casos de uso; modelos Eloquent para datos y relaciones; servicios para correo, seguridad y notificaciones; MySQL o MariaDB para persistencia; y Storage para archivos.')

doc.add_heading('2. Arquitectura de componentes',1)
table(doc,['Capa','Componentes principales','Responsabilidad'],[
 ('Cliente','Blade, Tailwind CSS, Alpine.js, Axios, Vite','Formularios, paneles, carga de evidencias, listados y experiencia de usuario.'),
 ('Web / HTTP','routes/web.php, controladores, middleware','Rutas, validación de entrada, autorización y coordinación de casos de uso.'),
 ('Dominio','Modelos Eloquent y relaciones','Reembolsos, centros de costos, aprobación, usuarios, eventos, fondos y reglas de acceso.'),
 ('Aplicación','Servicios, notificaciones y mailables','Correo Microsoft Graph/SMTP, agrupación de notificaciones, seguridad y preparación de adjuntos.'),
 ('Persistencia','Migraciones, seeders, Eloquent, MySQL/MariaDB','Esquema, datos de catálogo, permisos, historial y transacciones.'),
 ('Infraestructura','config, Storage, jobs, logs, public','Configuración, archivos, cola de base de datos, bitácoras y publicación web.')],[1500,3600,4260])
doc.add_heading('2.1 Flujo de solicitud',2)
table(doc,['Origen','Procesamiento','Destino'],[
 ('Usuario en navegador','Blade/Axios -> rutas web -> ReimbursementController','Modelos, Storage y respuestas de interfaz.'),
 ('XML/PDF/ticket','Validación y extracción de CFDI; preparación de adjuntos','Rutas de Storage y metadatos del reembolso.'),
 ('Decisión de aprobador','Reglas de paso, sustitución, perfil y centro de costos','reimbursement_approvals, paso actual y notificaciones.'),
 ('Cierre operativo','CXP Revisador y CXP Pagador/Tesorería','Semana de pago, estado final y exportaciones.')],[1800,4500,3060])

doc.add_heading('3. Tecnologías y dependencias',1)
table(doc,['Área','Tecnología implementada'],[
 ('Backend','PHP 8.2+, Laravel 12, Eloquent ORM.'),
 ('Frontend','Blade, Tailwind CSS 3, Alpine.js, Axios y Vite.'),
 ('Autenticación','Laravel Breeze, sesiones web, restablecimiento de contraseña, verificación de correo e invitaciones.'),
 ('Correo','Mailables y Notifications de Laravel; Microsoft Graph y configuración SMTP según ambiente.'),
 ('Documentos','Dompdf, FPDF, FPDI, PDF Parser y ZipStream para PDF, visualización y paquetes de archivos.'),
 ('Pruebas','PHPUnit y Laravel Test Runner; pruebas de autorización, flujo, archivos, CFDI y notificaciones.'),
 ('Colas','Conexión de base de datos por defecto; jobs, batches y tabla de fallos configurables.')],[2200,7160])

doc.add_heading('4. Módulos y responsabilidades',1)
table(doc,['Módulo','Puntos de entrada','Responsabilidad técnica'],[
 ('Reembolsos','ReimbursementController; recursos reimbursements','Alta normal y masiva, borradores, parseo CFDI, aprobación, aclaración, pago, archivos, exportaciones y auditoría.'),
 ('Usuarios y acceso','UserController, ProfileManagementController, middleware','Invitación, perfiles, permisos, sustitutos, usuarios activos y control de cuenta.'),
 ('Centros de costo','CostCenterController, ApprovalStep, BudgetRenewal','Presupuesto, responsables, categorías, matriz de aprobadores y correo Menfis/DMF.'),
 ('Eventos y fondos','TravelEventController, TravelEvent, FixedFund','Agrupación de gastos, presupuestos de evento y fondos fijos.'),
 ('Notificaciones','NotificationBatchService, NotificationBatch, Notifications','Notificaciones internas y por correo, agrupación y recordatorios.'),
 ('Seguridad de acceso','DeviceLoginService, LoginSecurityChallengeService, AccountBlockService','Trazabilidad de dispositivos, retos de acceso y bloqueo de cuentas.')],[2050,3120,4190])

doc.add_heading('5. Modelo de datos',1)
doc.add_paragraph('Las entidades centrales son users, profiles, permissions, cost_centers, approval_steps, reimbursements, reimbursement_files, reimbursement_approvals, travel_events, fixed_funds, notification_batches y sus tablas de relación. La integridad del proceso se apoya en relaciones Eloquent y en migraciones versionadas.')
table(doc,['Entidad','Relaciones y propósito'],[
 ('users','Propietario, creador, beneficiario, aprobador, sustituto y miembro autorizado de centros/eventos.'),
 ('cost_centers','Contiene presupuesto, responsables, usuarios autorizados, pasos de aprobación y correo de integración.'),
 ('approval_steps','Define el orden, aprobador y centro de costos del flujo dinámico.'),
 ('reimbursements','Agregado principal: tipo, estado, centro de costos, paso actual, datos fiscales y semana de pago.'),
 ('reimbursement_files','Evidencias XML, PDF, ticket y nombres originales.'),
 ('reimbursement_approvals','Bitácora de decisiones, comentarios, fechas, paso y sustituciones.'),
 ('notification_batches','Agrupa avisos para reducir mensajes individuales.')],[2200,7160])

doc.add_heading('6. Seguridad y control de acceso',1)
for t in [
 'Guard web basado en sesiones y proveedor Eloquent para usuarios.',
 'Middleware de autenticación, usuario activo, seguimiento de dispositivos y aliases role, permission y admin.',
 'RBAC mediante perfiles, permisos y asignación al centro de costos; el paso actual restringe quién puede aprobar.',
 'Protección CSRF y validación de formularios de Laravel; la solicitud de aclaración incluye límite de cinco intentos por minuto.',
 'Usuarios con soft delete, bloqueo de cuenta y registro de señales de dispositivos.',
 'Datos bancarios y detalles sensibles de compañía cuentan con migraciones de cifrado; los secretos permanecen en variables de entorno.']:
    doc.add_paragraph(t,style='List Bullet')

doc.add_heading('7. Integraciones, archivos y mensajería',1)
table(doc,['Integración','Diseño actual','Control relevante'],[
 ('CFDI','El controlador interpreta XML y conserva UUID, RFC, totales, impuestos y otros campos fiscales.','Revisar ante cambios de CFDI/SAT.'),
 ('Archivos','Laravel Storage conserva XML, PDF, ticket y archivos derivados.','Borradores inactivos por más de 30 días se depuran semanalmente junto con sus archivos.'),
 ('DMF/Menfis','El centro de costos define menfis_email; MenfisInvoiceMail envía XML y PDF cuando se cumplen precondiciones.','El nombre DMF no aparece literalmente; validar contra autorizaciones formales.'),
 ('Correo interno','Notifications y NotificationBatchService informan a aprobadores, solicitantes y colas CXP.','Evita saturación mediante agrupación.'),
 ('Microsoft','GraphMailService y transporte Microsoft Graph amplían el envío de correo.','Credenciales y configuración fuera del repositorio, en ambiente.')],[2000,4700,2660])

doc.add_heading('8. Operación y despliegue',1)
doc.add_paragraph('El despliegue requiere configurar variables de aplicación, base de datos, almacenamiento y correo; instalar dependencias PHP y JavaScript; generar la clave; ejecutar migraciones; compilar recursos Vite y apuntar el servidor web al directorio public.')
table(doc,['Componente','Consideración operativa'],[
 ('Base de datos','Aplicar migraciones controladas y respaldar antes de cambios estructurales.'),
 ('Colas','La conexión predeterminada es database. Ejecutar worker en ambientes que procesen trabajos; monitorear failed_jobs.'),
 ('Tareas programadas','Se elimina semanalmente el contenido de borradores inactivos por más de 30 días. También existen comandos para recordatorios y procesamiento de notificaciones.'),
 ('Archivos','Garantizar permisos de escritura en storage y bootstrap/cache; proteger rutas de descarga mediante autorización.'),
 ('Observabilidad','Revisar logs de Laravel, errores de adjuntos, fallos de correo y eventos de seguridad.')],[2300,7060])

doc.add_heading('9. Riesgos y evolución recomendada',1)
table(doc,['Prioridad','Recomendación'],[
 ('Alta','Formalizar la máquina de estados en un servicio o enum de dominio y reducir el uso de estados históricos en pantallas y reportes.'),
 ('Alta','Versionar las autorizaciones y destinatarios de DMF/Menfis, incluyendo evidencia, vigencia y aprobador.'),
 ('Media','Mover los envíos externos a cola con reintentos, trazabilidad de destinatario/UUID/resultado y alertas de fallos.'),
 ('Media','Mantener pruebas para autoaprobación, sustitución, corrección, CXP, pago, archivos y correo.'),
 ('Media','Aislar la lógica de parseo CFDI para facilitar actualizaciones ante cambios del SAT.'),
 ('Baja','Documentar objetivos de recuperación, retención de archivos y monitoreo de capacidad.')],[1300,8060])

doc.add_heading('10. Referencias de código',1)
for t in ['composer.json y package.json.', 'bootstrap/app.php, config/auth.php y config/queue.php.', 'routes/web.php.', 'app/Http/Controllers/ReimbursementController.php.', 'app/Models/Reimbursement.php, CostCenter.php y User.php.', 'app/Services/*, app/Notifications/*, app/Mail/* y database/migrations/*.']:
    doc.add_paragraph(t,style='List Bullet')
doc.save(OUT); print(OUT)
