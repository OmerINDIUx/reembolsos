from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.section import WD_SECTION

OUT = 'Documento_Funcional_Flujos_DMF.docx'
BLUE = '2E74B5'
DARK = '1F4D78'
INK = '203040'
PALE = 'E8EEF5'
GRAY = 'F2F4F7'

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), fill); tcPr.append(shd)

def margins(cell, top=80, start=120, bottom=80, end=120):
    tcPr = cell._tc.get_or_add_tcPr()
    mar = tcPr.first_child_found_in('w:tcMar')
    if mar is None:
        mar = OxmlElement('w:tcMar'); tcPr.append(mar)
    for key, val in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = mar.find(qn('w:' + key))
        if node is None:
            node = OxmlElement('w:' + key); mar.append(node)
        node.set(qn('w:w'), str(val)); node.set(qn('w:type'), 'dxa')

def set_widths(table, widths):
    table.autofit = False
    tblPr = table._tbl.tblPr
    tblW = tblPr.first_child_found_in('w:tblW')
    if tblW is None:
        tblW = OxmlElement('w:tblW'); tblPr.append(tblW)
    tblW.set(qn('w:w'), str(sum(widths))); tblW.set(qn('w:type'), 'dxa')
    ind = OxmlElement('w:tblInd'); ind.set(qn('w:w'), '120'); ind.set(qn('w:type'), 'dxa'); tblPr.append(ind)
    grid = table._tbl.tblGrid
    for col, w in zip(grid.gridCol_lst, widths): col.set(qn('w:w'), str(w))
    for row in table.rows:
        for cell, w in zip(row.cells, widths):
            cell.width = Inches(w / 1440)
            tcW = cell._tc.tcPr.tcW
            tcW.set(qn('w:w'), str(w)); tcW.set(qn('w:type'), 'dxa')
            margins(cell); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

def para(doc, text='', style=None, bold_prefix=None, italic=False):
    p = doc.add_paragraph(style=style)
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix); r.bold = True
        p.add_run(text[len(bold_prefix):])
    else:
        r = p.add_run(text); r.italic = italic
    return p

def heading(doc, text, level=1):
    return doc.add_paragraph(text, style=f'Heading {level}')

def table(doc, headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    t.style = 'Table Grid'
    for cell, value in zip(t.rows[0].cells, headers):
        shade(cell, PALE)
        p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(0)
        r = p.add_run(value); r.bold = True; r.font.size = Pt(9); r.font.color.rgb = RGBColor.from_string(DARK)
    for row in rows:
        cells = t.add_row().cells
        for cell, value in zip(cells, row):
            p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(0)
            r = p.add_run(value); r.font.size = Pt(9)
    set_widths(t, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t

doc = Document()
sec = doc.sections[0]
sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Inches(1)
sec.header_distance = sec.footer_distance = Inches(.492)

styles = doc.styles
normal = styles['Normal']; normal.font.name = 'Calibri'; normal.font.size = Pt(11); normal.font.color.rgb = RGBColor.from_string(INK)
normal.paragraph_format.space_after = Pt(6); normal.paragraph_format.line_spacing = 1.10
for level, size, color, before, after in [(1,16,BLUE,16,8),(2,13,BLUE,12,6),(3,12,DARK,8,4)]:
    st = styles[f'Heading {level}']; st.font.name='Calibri'; st.font.size=Pt(size); st.font.bold=True; st.font.color.rgb=RGBColor.from_string(color)
    st.paragraph_format.space_before=Pt(before); st.paragraph_format.space_after=Pt(after); st.paragraph_format.keep_with_next=True

header = sec.header.paragraphs[0]
header.text = 'PLATAFORMA DE REEMBOLSOS  |  GUÍA OPERATIVA'
header.runs[0].font.size = Pt(8); header.runs[0].font.color.rgb = RGBColor.from_string('6B7280')
header.paragraph_format.space_after = Pt(0)
footer = sec.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('Documento de referencia interna  |  3 de septiembre de 2026').font.size = Pt(8)

p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before = Pt(42); p.paragraph_format.space_after = Pt(8)
r = p.add_run('PLATAFORMA DE REEMBOLSOS'); r.bold=True; r.font.size=Pt(11); r.font.color.rgb=RGBColor.from_string(BLUE)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(12)
r = p.add_run('Roles, flujos, reglas y matriz de estados'); r.bold=True; r.font.size=Pt(24); r.font.color.rgb=RGBColor.from_string(DARK)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(26)
r = p.add_run('Documento funcional y de control operativo'); r.italic=True; r.font.size=Pt(12); r.font.color.rgb=RGBColor.from_string('5B6573')

table(doc, ['Control documental', 'Valor'], [
    ('Versión', '1.0'), ('Fecha', '3 de septiembre de 2026'), ('Fuente verificable', 'Implementación vigente del repositorio reembolsos'),
    ('Destinatarios', 'Operación, Cuentas por Pagar, Tesorería, Administración y Tecnología')], [2700, 6660])
p = doc.add_paragraph(); shade_cell = doc.add_table(rows=1, cols=1).cell(0,0); shade(shade_cell, GRAY); margins(shade_cell, 120, 160, 120, 160)
sp = shade_cell.paragraphs[0]; sp.paragraph_format.space_after=Pt(0)
rr = sp.add_run('Nota de evidencia. '); rr.bold=True; rr.font.color.rgb=RGBColor.from_string(DARK)
sp.add_run('El repositorio no contiene los correos originales de DMF. Por ello, el apartado DMF/Menfis describe el comportamiento implementado y señala explícitamente los puntos que requieren cotejo con los correos de autorización.').font.size=Pt(10)

heading(doc, '1. Propósito y alcance')
para(doc, 'La plataforma administra solicitudes de reembolso y comprobación de gastos desde su captura hasta el pago. Integra evidencias fiscales, centros de costos, aprobación multinivel, revisión de Cuentas por Pagar, autorización de pago, notificaciones y auditoría.')
para(doc, 'El alcance cubre reposiciones, viáticos, comprobaciones de tarjeta empresarial, fondos fijos y gastos asociados a eventos o viajes. Las evidencias principales son XML de CFDI, PDF de factura, ticket y documentación adicional.')

heading(doc, '2. Roles y responsabilidades')
table(doc, ['Rol', 'Responsabilidad', 'Alcance'], [
 ('Administrador', 'Configura usuarios, perfiles, centros de costos, permisos y excepciones.', 'Global, conforme a permisos administrativos.'),
 ('Administrador de lectura', 'Consulta, auditoría y seguimiento.', 'Sin cambios operativos.'),
 ('Solicitante', 'Crea borradores, adjunta evidencias, envía y atiende correcciones.', 'Solicitudes propias o creadas por terceros autorizados.'),
 ('Director / Control de Obra / Director Ejecutivo', 'Aprueba los pasos operativos configurados.', 'Solicitudes del centro de costos y paso asignado.'),
 ('CXP Revisador', 'Revisa la documentación posterior al flujo operativo.', 'Cola pendiente_revision_cxp.'),
 ('Subdirección', 'Autoriza el nivel directivo cuando está configurado.', 'Solicitudes asignadas por centro de costos.'),
 ('CXP Pagador / Tesorería', 'Autoriza pago y registra semana de procesamiento.', 'Cola pendiente_pago.')], [2250, 4400, 2710])
para(doc, 'Modelo de autorización: los permisos del perfil (RBAC), la asignación por centro de costos y los pasos dinámicos de aprobación se combinan. Se permiten sustituciones temporales; la bitácora conserva tanto al sustituto como al titular asignado cuando aplica.', italic=True)

heading(doc, '3. Flujo operativo autorizado')
heading(doc, '3.1 Secuencia principal', 2)
table(doc, ['Etapa', 'Acción y control', 'Resultado'], [
 ('1. Captura', 'El solicitante crea borrador y adjunta XML, PDF y/o ticket. El sistema analiza y conserva datos del CFDI.', 'Borrador listo para envío.'),
 ('2. Envío', 'Se identifica el centro de costos y su flujo ordenado. Si hay pasos, se asigna el primero; si no, puede autoaprobarse.', 'Pendiente de aprobación o revisión CXP.'),
 ('3. Aprobación', 'El asignado, sustituto o administrador decide. Se registra paso, fecha, comentario y sustitución.', 'Siguiente paso, CXP, corrección o rechazo.'),
 ('4. Revisión CXP', 'CXP Revisador valida documentación.', 'Pendiente de pago.'),
 ('5. Pago', 'CXP Pagador/Tesorería autoriza y asigna semana de pago.', 'Pagado al confirmar el pago.')], [1500, 5400, 2460])
heading(doc, '3.2 Excepciones y controles', 2)
para(doc, 'Si el solicitante coincide con aprobadores de niveles inferiores o iguales, esos pasos pueden autoaprobarse y deben quedar justificados en la bitácora. Director Ejecutivo y Subdirección inician normalmente en Control de Obra cuando ese nivel existe. Si una persona ya aprobó el documento en una etapa previa, el siguiente paso puede autoaprobarse.')
para(doc, 'La solicitud puede volver a corrección sin perder trazabilidad. Un rechazo finaliza el flujo. Toda decisión queda registrada en reimbursement_approvals.')

heading(doc, '4. Matriz de estados')
table(doc, ['Estado', 'Significado', 'Actor', 'Transición típica'], [
 ('borrador', 'Captura incompleta guardada.', 'Solicitante / creador.', 'pendiente o en_evento.'),
 ('en_evento', 'Gasto asociado a evento aún no enviado.', 'Participante / responsable.', 'pendiente.'),
 ('pendiente', 'Esperando el paso actual.', 'Aprobador asignado / sustituto / administrador.', 'pendiente, revisión CXP, corrección o rechazo.'),
 ('pendiente_revision_cxp', 'En espera de revisión documental/contable.', 'CXP Revisador.', 'pendiente_pago.'),
 ('pendiente_pago', 'En espera de autorización de pago.', 'CXP Pagador / Tesorería.', 'pagado.'),
 ('pagado', 'Pago confirmado.', 'CXP Pagador / Tesorería.', 'Fin.'),
 ('requiere_correccion', 'Faltan datos o evidencias.', 'Solicitante / creador.', 'pendiente o rechazado.'),
 ('rechazado', 'Solicitud no autorizada.', 'Aprobador autorizado.', 'Fin.'),
 ('aprobado_* / aprobado', 'Marcas históricas para compatibilidad y reportes.', 'Sistema / auditoría.', 'Según contexto financiero.')], [2100, 3050, 1960, 2250])
para(doc, 'Nota: el flujo vigente se guía principalmente por current_step_id, approval_steps, pendiente_revision_cxp y pendiente_pago. Los estados aprobado_* se conservan para compatibilidad histórica.', italic=True)

heading(doc, '5. Reglas de negocio')
heading(doc, '5.1 Autorización y flujo', 2)
for txt in [
 'Un aprobador sólo actúa sobre su paso actual, salvo sustitución activa o privilegio administrativo.',
 'Los pasos se ejecutan en el orden definido en approval_steps.order.',
 'Al terminar los pasos operativos, la solicitud pasa a pendiente_revision_cxp antes de pago.',
 'CXP Revisador actúa únicamente en pendiente_revision_cxp; CXP Pagador/Tesorería, en pendiente_pago.',
 'Los usuarios consultan solicitudes según perfil, relación con el centro de costos, paso asignado o pertenencia a las colas CXP.']:
    para(doc, txt, style='List Bullet')
heading(doc, '5.2 Evidencia, fiscalidad y presupuesto', 2)
for txt in [
 'Una factura con XML conserva UUID y datos críticos del CFDI para auditoría y consulta.',
 'XML y PDF son las evidencias principales de la integración de correo DMF/Menfis.',
 'Los archivos de borrador tienen un límite de carga de 10 MB.',
 'La solicitud se vincula a un centro de costos activo y opcionalmente a un fondo fijo o evento; la semana de pago se registra al autorizar el pago.']:
    para(doc, txt, style='List Bullet')

heading(doc, '6. Flujo de correo DMF/Menfis')
para(doc, 'El destinatario se configura por centro de costos mediante menfis_email. El código no utiliza literalmente el nombre DMF; para este documento se trata como el buzón DMF/Menfis sujeto a validación formal.')
table(doc, ['Control', 'Comportamiento implementado'], [
 ('Disparador', 'Alta normal, carga masiva o reenvío posterior a corrección, siempre que exista UUID válido.'),
 ('Precondición', 'El centro de costos debe tener menfis_email con formato válido.'),
 ('Destinatario', 'menfis_email configurado en el centro de costos.'),
 ('Asunto', 'Nombre original del XML sin extensión; si no existe, factura.'),
 ('Adjuntos', 'XML original y PDF asociado; el PDF toma la base del nombre XML cuando está disponible.'),
 ('Sin envío', 'No se envía si falta un correo válido configurado.'),
 ('Tolerancia a falla', 'Errores de adjuntos se registran en logs sin interrumpir el flujo principal.')], [2100, 7260])
heading(doc, '6.1 Pendientes para acreditar la autorización por correo', 2)
para(doc, 'Para convertir el flujo anterior en una transcripción formal de autorizaciones DMF, deben anexarse los correos fuente con: remitente, fecha, destinatarios, asunto, condiciones de envío, excepciones y texto aprobado. La fuente verificable disponible hoy es la implementación de MenfisInvoiceMail, CostCenter::menfisEmailAddress() y los puntos de envío del controlador.')

heading(doc, '7. Trazabilidad, riesgos y acciones recomendadas')
table(doc, ['Tema', 'Control / recomendación'], [
 ('Trazabilidad', 'Conservar propietario, creador, beneficiario, paso actual, decisiones, observaciones, sustituciones, evidencias y datos fiscales. reimbursement_approvals es la bitácora funcional.'),
 ('Riesgo DMF', 'El destinatario se configura por centro de costos y no existe una lista de autorización versionada en la base de datos.'),
 ('Acción prioritaria', 'Incorporar los correos formales DMF como anexo controlado con fecha, versión y aprobador.'),
 ('Acción técnica', 'Registrar destinatario, fecha, UUID, resultado y reintentos del envío; mover a cola si aumenta el volumen.'),
 ('Calidad de flujo', 'Probar autoaprobación, sustitución, corrección, CXP, pago y envío DMF/Menfis; consolidar los estados históricos gradualmente.')], [2500, 6860])

heading(doc, '8. Fuentes de implementación')
for txt in [
 'app/Models/User.php; app/Models/Reimbursement.php; app/Models/CostCenter.php.',
 'app/Http/Controllers/ReimbursementController.php y app/Mail/MenfisInvoiceMail.php.',
 'app/Notifications/*, database/migrations/*, project_architecture.md y DOCUMENTACION_FUNCIONAL_TECNICA.md.']:
    para(doc, txt, style='List Bullet')

doc.save(OUT)
print(OUT)
